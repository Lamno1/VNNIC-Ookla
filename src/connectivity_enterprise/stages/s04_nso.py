from pathlib import Path
import re
import unicodedata
import json

import openpyxl
import pandas as pd

from ..core.contracts import require_unique
from ..core.paths import WORKING_ROOT, sources_config


TABLES = {
    "V05.02.xlsx": "new_registrations",
    "V05.04.xlsx": "active_firms_dec31",
    "V05.05.xlsx": "active_firms_per_1000",
}


def norm(value):
    value = str(value or "").strip().lower().replace("đ", "d")
    value = unicodedata.normalize("NFD", value)
    value = "".join(c for c in value if unicodedata.category(c) != "Mn")
    value = re.sub(r"\btp\.?\s*", "", value)
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return " ".join(value.split())


def run(ctx):
    nso_root = Path(sources_config()["protected_roots"]["nso"])
    dim = pd.read_parquet(WORKING_ROOT / "data/reference/dim_province_legacy63.parquet")
    official_lookup = {
        norm(row.province_name_official): row
        for row in dim.itertuples(index=False)
    }

    facts = []
    exclusions = []
    matches = []

    for filename, measure_id in TABLES.items():
        path = nso_root / filename
        wb = openpyxl.load_workbook(
            path, read_only=True, data_only=True, keep_links=False
        )
        if len(wb.sheetnames) != 1:
            raise RuntimeError(f"FAIL-CLOSED: {filename} has unexpected sheets")
        ws = wb[wb.sheetnames[0]]
        title = ws.cell(1, 1).value
        years = [int(ws.cell(3, col).value) for col in range(2, ws.max_column + 1)]
        table_matches = 0
        table_exclusions = 0

        for row_no in range(4, ws.max_row + 1):
            label = ws.cell(row_no, 1).value
            normalized = norm(label)
            dim_row = official_lookup.get(normalized)

            if dim_row is None:
                exclusions.append({
                    "source_workbook": filename,
                    "source_sheet": ws.title,
                    "source_row": row_no,
                    "original_label": label,
                    "normalized_label": normalized,
                    "exclusion_status": "EXCLUDED_FROM_PROVINCE_CANDIDATE",
                    "exclusion_reason": "NATIONAL_REGION_OR_UNRESOLVED_AGGREGATE",
                })
                table_exclusions += 1
                continue

            matches.append({
                "source_workbook": filename,
                "source_sheet": ws.title,
                "source_row": row_no,
                "source_label": label,
                "normalized_label": normalized,
                "province_id_legacy63": dim_row.province_id_legacy63,
                "province_name_official": dim_row.province_name_official,
                "match_method": "normalized_name_exact_candidate",
                "review_status": "PENDING_AUTHORITATIVE_IDENTIFIER_REVIEW",
            })
            table_matches += 1
            for col, year in enumerate(years, start=2):
                value = ws.cell(row_no, col).value
                facts.append({
                    "province_id_legacy63": dim_row.province_id_legacy63,
                    "year": year,
                    "measure_id": measure_id,
                    "value": value,
                    "source_workbook": str(path),
                    "source_sheet": ws.title,
                    "source_row_or_range": f"{ws.title}!{ws.cell(row_no, col).coordinate}",
                    "original_label": title,
                    "unit": "",
                    "price_basis": "",
                    "semantic_status": "UNVERIFIED",
                    "exclusion_status": "CANDIDATE_NOT_RELEASED",
                    "exclusion_reason": "G1_G2_G5_PENDING",
                })

        if table_matches != 63 or table_exclusions != 7:
            raise RuntimeError(
                f"FAIL-CLOSED: {filename} expected 63 province candidates "
                f"and 7 aggregate exclusions; found {table_matches} and {table_exclusions}"
            )

    fact = pd.DataFrame(facts).sort_values(
        ["measure_id", "province_id_legacy63", "year"]
    ).reset_index(drop=True)
    require_unique(fact, ["province_id_legacy63", "year", "measure_id"], "nso candidate fact")

    expected_rows = 63 * (9 + 8 + 8)
    if len(fact) != expected_rows:
        raise RuntimeError(
            f"FAIL-CLOSED: expected {expected_rows} NSO candidate cells, found {len(fact)}"
        )

    fact_path = WORKING_ROOT / "data/quarantine/fact_nso_measure_candidates.parquet"
    fact.to_parquet(fact_path, index=False)
    matches_path = WORKING_ROOT / "artifacts/qc/nso_name_crosswalk_review.csv"
    pd.DataFrame(matches).sort_values(
        ["source_workbook", "province_id_legacy63"]
    ).to_csv(matches_path, index=False, encoding="utf-8-sig")
    exclusions_path = WORKING_ROOT / "artifacts/qc/nso_row_exclusions.csv"
    pd.DataFrame(exclusions).sort_values(
        ["source_workbook", "source_row"]
    ).to_csv(exclusions_path, index=False, encoding="utf-8-sig")

    qc = {
        "status": "FAIL_PENDING_PROVENANCE_SEMANTICS_AND_IDENTIFIER_REVIEW",
        "workbooks": 3,
        "candidate_fact_rows": len(fact),
        "province_candidates_per_table": 63,
        "aggregate_exclusions_per_table": 7,
        "missing_numeric_cells": int(pd.to_numeric(fact["value"], errors="coerce").isna().sum()),
        "name_mapping_review_status": "PENDING",
        "semantic_status": "UNVERIFIED",
        "released_fact_nso_measure": False,
        "released_fact_enterprise_outcomes": False,
    }
    qc_path = WORKING_ROOT / "artifacts/qc/nso_structural_audit.json"
    qc_path.write_text(json.dumps(qc, indent=2), encoding="utf-8")

    official_decision_path = WORKING_ROOT / "artifacts/qc/nso_semantic_decision.json"
    official_decision = json.loads(official_decision_path.read_text(encoding="utf-8")) if official_decision_path.exists() else {}
    official_pass = official_decision.get("G5_PRIMARY_OUTCOME") == "PASS"
    qc["status"] = "PASS_OFFICIAL_SNAPSHOT_PRIMARY_LOCAL_DERIVATIVES_QUARANTINED" if official_pass else qc["status"]
    if official_pass:
        qc["name_mapping_review_status"] = "LOCAL_WORKBOOK_MAPPING_NOT_USED"
        qc["semantic_status"] = official_decision.get("G2_NSO_PRIMARY")
        qc["official_primary_source_available"] = True
        qc["local_candidate_status"] = "QUARANTINED_NOT_USED"
    qc["official_snapshot_primary_outcome"] = official_decision.get("G2_NSO_PRIMARY", "NOT_AVAILABLE")
    qc["local_candidates_remain_quarantined"] = True
    qc_path.write_text(json.dumps(qc, indent=2), encoding="utf-8")
    ctx["gatebook"].set(
        "G5", "PASS" if official_pass else "FAIL",
        "Official PXWeb V05.05 snapshot and deterministic 63-code audit pass; local workbook candidates remain quarantined."
        if official_pass else "NSO official primary-outcome evidence is insufficient; candidates remain quarantined.",
        {**qc, "official_decision": str(official_decision_path)},
    )
    outputs = [fact_path, matches_path, exclusions_path, qc_path]
    for evidence_path in [
        WORKING_ROOT / "artifacts/qc/nso_official_reconciliation.csv",
        WORKING_ROOT / "artifacts/qc/province_identifier_audit.csv",
        WORKING_ROOT / "artifacts/manifests/nso_official_snapshot_manifest.csv",
        WORKING_ROOT / "artifacts/lineage/semantic_evidence_lineage.csv",
    ]:
        if evidence_path.exists():
            outputs.append(evidence_path)
    return outputs
